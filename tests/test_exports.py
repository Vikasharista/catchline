from app.exports.award_xlsx import build_award_xlsx
from app.exports.memo import build_memo_pdf
from app.exports.rfq_pdf import build_rfq_pdf
from app.exports.template_xlsx import build_template_xlsx

DRAFT = {
    "scope": {
        "title": "Frozen seafood RFx",
        "incoterm": "DAP",
        "incoterm_place": "Boulogne",
        "currency": "EUR",
        "weight_basis": "net",
        "response_deadline": "2026-09-10",
        "award_date": "2026-09-30",
        "contract_start": "2026-10-01",
        "contract_end": "2027-09-30",
    },
    "lines": [
        {"line_id": "L01", "species": "Atlantic salmon", "form": "HOG", "grade": "2-3 kg", "annual_volume_kg": 40000}
    ],
    "questionnaire": [
        {"q_id": "Q01", "group": "certifications", "text": "Certificate?", "answer_type": "text"}
    ],
}


def test_build_rfq_pdf_produces_pdf_bytes():
    pdf_bytes = build_rfq_pdf(DRAFT)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(pdf_bytes) > 500


def test_build_template_xlsx_has_quote_and_questionnaire_sheets():
    xlsx_bytes = build_template_xlsx(DRAFT)
    assert xlsx_bytes[:2] == b"PK"  # xlsx is a zip
    assert len(xlsx_bytes) > 500


def test_build_award_xlsx_has_all_tabs():
    from openpyxl import load_workbook
    import io

    xlsx_bytes = build_award_xlsx(
        award_table=[{"line_id": "L01", "supplier_id": "S1", "eur_kg_net_dap": 7.25, "volume_kg": 40000}],
        comparison_rows=[{"line_id": "L01", "supplier_id": "S1", "eur_kg_net_dap": 7.25, "status": "confirmed", "eligibility": "eligible"}],
        assumptions={"fx_usd": 1.17},
        open_flags=[{"supplier_id": "S4", "line_id": "L28", "code": "illegible", "severity": "red", "message": "stained"}],
        eligibility=[{"supplier_id": "S4", "status": "not_eligible", "blocked_lines": [], "reasons": ["expired cert"]}],
        audit_log=[{"ts": "2026-09-10", "actor": "buyer", "action": "sent", "entity": "rfx:1", "reason": None}],
    )
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert set(wb.sheetnames) == {"Award", "Comparison", "Assumptions", "Open flags", "Eligibility", "Audit log"}
    assert wb["Award"]["A2"].value == "L01"


def test_build_memo_pdf_produces_pdf_bytes():
    pdf_bytes = build_memo_pdf(
        recommendation_text="Award L01 to Fjordline at EUR 7.25/kg.",
        total_spend_eur=290000.0,
        award_table=[{"line_id": "L01", "supplier_id": "S1", "eur_kg_net_dap": 7.25}],
        risks=["Baltic Blue certificate pending"],
        assumptions={"fx_usd": 1.17},
        n_logged_decisions=12,
    )
    assert pdf_bytes[:4] == b"%PDF"
    assert len(pdf_bytes) > 500
