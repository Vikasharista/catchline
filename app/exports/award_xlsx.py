"""Builds award.xlsx (PRD §7.10): Award, Comparison, Assumptions, Open
flags, Eligibility, Audit log tabs — all straight from tool results and the
DB, no LLM involved.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)


def build_award_xlsx(
    award_table: list[dict],
    comparison_rows: list[dict],
    assumptions: dict,
    open_flags: list[dict],
    eligibility: list[dict],
    audit_log: list[dict],
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    _sheet(
        wb,
        "Award",
        ["Line", "Supplier", "EUR/kg net DAP", "Volume (kg)", "Line spend (EUR)"],
        [
            [r["line_id"], r["supplier_id"], r["eur_kg_net_dap"], r.get("volume_kg"), round(r["eur_kg_net_dap"] * r.get("volume_kg", 0), 2)]
            for r in award_table
        ],
    )

    _sheet(
        wb,
        "Comparison",
        ["Line", "Supplier", "EUR/kg net DAP", "Status", "Eligibility"],
        [
            [r.get("line_id"), r.get("supplier_id"), r.get("eur_kg_net_dap"), r.get("status"), r.get("eligibility")]
            for r in comparison_rows
        ],
    )

    _sheet(wb, "Assumptions", ["Key", "Value"], [[k, str(v)] for k, v in assumptions.items()])

    _sheet(
        wb,
        "Open flags",
        ["Supplier", "Line", "Code", "Severity", "Message"],
        [[f.get("supplier_id"), f.get("line_id"), f.get("code"), f.get("severity"), f.get("message")] for f in open_flags],
    )

    _sheet(
        wb,
        "Eligibility",
        ["Supplier", "Status", "Blocked lines", "Reasons"],
        [
            [e.get("supplier_id"), e.get("status"), ", ".join(e.get("blocked_lines", [])), "; ".join(e.get("reasons", []))]
            for e in eligibility
        ],
    )

    _sheet(
        wb,
        "Audit log",
        ["Timestamp", "Actor", "Action", "Entity", "Reason"],
        [[a.get("ts"), a.get("actor"), a.get("action"), a.get("entity"), a.get("reason")] for a in audit_log],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
