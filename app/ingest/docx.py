from pathlib import Path

import docx

from app.ingest.router import Chunk, IngestedDoc


def ingest(path: Path) -> IngestedDoc:
    document = docx.Document(path)
    chunks: list[Chunk] = []

    body_children = document.element.body
    para_idx = 0
    table_idx = 0
    # Walk the document body in order so paragraph/table locators reflect
    # reading order, not just python-docx's separate .paragraphs/.tables lists.
    para_iter = iter(document.paragraphs)
    table_iter = iter(document.tables)

    for child in body_children:
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = next(para_iter, None)
            para_idx += 1
            if para is not None and para.text.strip():
                chunks.append(Chunk(kind="text", content=para.text, locator=f"paragraph {para_idx}"))
        elif tag == "tbl":
            table = next(table_iter, None)
            table_idx += 1
            if table is not None:
                rows = ["\t".join(cell.text for cell in row.cells) for row in table.rows]
                chunks.append(Chunk(kind="table", content="\n".join(rows), locator=f"table {table_idx}"))

    return IngestedDoc(filename=path.name, kind="docx", chunks=chunks)
